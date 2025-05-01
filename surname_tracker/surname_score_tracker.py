import openpyxl
import os

EXCEL_FILE = "student_scores.xlsx"

# Step 1: Create or load the workbook and sheet
def create_or_load_workbook():
    if not os.path.exists(EXCEL_FILE):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "Scores"
        sheet.append(["Student Name", "Score", "Status"])
        workbook.save(EXCEL_FILE)
    else:
        workbook = openpyxl.load_workbook(EXCEL_FILE)
    return workbook

# Step 2: Add student data and calculate pass/fail
def add_student_score(name, score):
    workbook = create_or_load_workbook()
    sheet = workbook["Scores"]
    status = "Pass" if score >= 75 else "Fail"
    sheet.append([name, score, status])
    workbook.save(EXCEL_FILE)
    print(f"Added: {name} - {score} ({status})")

# Step 4: Retrieve and display all records
def display_all_records():
    workbook = create_or_load_workbook()
    sheet = workbook["Scores"]
    print("\nAll Student Records:")
    for row in sheet.iter_rows(min_row=2, values_only=True):
        print(f"Name: {row[0]}, Score: {row[1]}, Status: {row[2]}")

# Main logic
def main():
    while True:
        print("\n--- Student Score Tracker ---")
        print("1. Add Student Score")
        print("2. Show All Records")
        print("3. Exit")
        choice = input("Enter choice (1/2/3): ")

        if choice == "1":
            name = input("Enter student name: ")
            try:
                score = int(input("Enter score (0-100): "))
                if 0 <= score <= 100:
                    add_student_score(name, score)
                else:
                    print("Score must be between 0 and 100.")
            except ValueError:
                print("Invalid input. Please enter a number.")
        elif choice == "2":
            display_all_records()
        elif choice == "3":
            print("Exiting program.")
            break
        else:
            print("Invalid choice. Try again.")

if __name__ == "__main__":
    main()
